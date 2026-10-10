from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime
import os


class WordGenerator:

    def __init__(self, output_dir="data/output"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    # =========================================================
    # CONFIGURATION WORD
    # =========================================================

    def setup_document(self, document):

        section = document.sections[0]

        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

        styles = document.styles

        normal = styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10)

        for style_name in ["Title", "Heading 1", "Heading 2"]:
            style = styles[style_name]
            style.font.name = "Aptos"

    # =========================================================
    # NUMEROTATION DES PAGES
    # =========================================================

    def add_page_number(self, section):

        footer = section.footer

        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run("Page ")

        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")

        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = "PAGE"

        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")

        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

    # =========================================================
    # TEXTE
    # =========================================================

    def add_text(self, document, text, bold=False):

        if not text:
            return

        paragraph = document.add_paragraph()

        run = paragraph.add_run(str(text))
        run.bold = bold
        run.font.size = Pt(10)

        paragraph.paragraph_format.space_after = Pt(5)

    def add_bullet(self, document, text):

        if not text:
            return

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        run = paragraph.add_run(str(text))
        run.font.size = Pt(10)

    # =========================================================
    # TITRE DE SECTION
    # =========================================================

    def add_section_title(self, document, title):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)

        run = paragraph.add_run(title)

        run.bold = True
        run.font.size = Pt(13)

    # =========================================================
    # PAGE DE GARDE
    # =========================================================

    def add_cover(self, document, number_items):

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.paragraph_format.space_before = Pt(100)

        run = paragraph.add_run(
            "SMARTLY KNOWLEDGE BASE"
        )

        run.bold = True
        run.font.size = Pt(24)

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            "MODE OPÉRATOIRE"
        )

        run.bold = True
        run.font.size = Pt(20)

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            "Prise en charge des clients par motif de contact"
        )

        run.font.size = Pt(15)

        document.add_paragraph()

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            f"{number_items} motifs de contact"
        )

        run.font.size = Pt(12)

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            f"Généré le {datetime.now().strftime('%d/%m/%Y')}"
        )

        run.font.size = Pt(10)

        document.add_page_break()

    # =========================================================
    # UN MOTIF MODOP
    # =========================================================

    def add_modop_item(
        self,
        document,
        index,
        row,
        metadata
    ):

        motif = row.get("motif", "")
        pourquoi = row.get("pourquoi", "")
        que_faire = row.get("que_faire", "")
        codification = row.get("codification", "")
        statut = row.get("statut", "")

        # -----------------------------------------------------
        # MOTIF
        # -----------------------------------------------------

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            f"{index}. {motif}"
        )

        run.bold = True
        run.font.size = Pt(15)

        # -----------------------------------------------------
        # INTENTION
        # -----------------------------------------------------

        self.add_section_title(
            document,
            "INTENTION CLIENT"
        )

        self.add_text(
            document,
            metadata["intention"]
        )

        # -----------------------------------------------------
        # POURQUOI
        # -----------------------------------------------------

        if pourquoi:

            self.add_section_title(
                document,
                "POURQUOI ?"
            )

            self.add_text(
                document,
                pourquoi
            )

        # -----------------------------------------------------
        # QUESTIONS UTILISATEUR
        # -----------------------------------------------------

        self.add_section_title(
            document,
            "QUESTIONS UTILISATEUR PROBABLES"
        )

        for question in metadata["questions"]:
            self.add_bullet(
                document,
                question
            )

        # -----------------------------------------------------
        # QUE FAIRE
        # -----------------------------------------------------

        if que_faire:

            self.add_section_title(
                document,
                "QUE FAIRE ?"
            )

            self.add_text(
                document,
                que_faire
            )

        # -----------------------------------------------------
        # CODIFICATION
        # -----------------------------------------------------

        if codification:

            self.add_section_title(
                document,
                "CODIFICATION"
            )

            self.add_text(
                document,
                codification
            )

        # -----------------------------------------------------
        # STATUT
        # -----------------------------------------------------

        if statut:

            self.add_section_title(
                document,
                "STATUT"
            )

            self.add_text(
                document,
                statut
            )

        # Séparateur
        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "────────────────────────────────────────"
        )

        run.font.size = Pt(8)

    # =========================================================
    # GENERATION MODOP
    # =========================================================

    def generate_modop_document(
        self,
        df,
        ai_metadata
    ):

        document = Document()

        self.setup_document(document)

        # Numéro de page
        self.add_page_number(
            document.sections[0]
        )

        # Couverture
        self.add_cover(
            document,
            len(df)
        )

        # Contenus
        for index, (_, row) in enumerate(
            df.iterrows(),
            start=1
        ):

            metadata = ai_metadata.generate_modop_metadata(
                row["motif"],
                row.get("pourquoi", ""),
                row.get("que_faire", "")
            )

            self.add_modop_item(
                document,
                index,
                row,
                metadata
            )

        output_file = os.path.join(
            self.output_dir,
            "MODOP_Prise_en_Charge_Client.docx"
        )

        document.save(output_file)

        return output_file